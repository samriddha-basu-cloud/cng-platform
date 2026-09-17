"""Tests for the self-generating Excel template/sample/import pipeline
(spec sections 37-57). The app must never require an external Excel file:
these tests exercise the exact round trip a real user would (download
sample -> upload it unmodified -> it validates clean -> it imports), then
probe the validation engine with deliberately broken rows.
"""
import io
from app.services import excel_service as xs


def test_template_has_all_required_sheets_and_readme():
    buf = xs.generate_template_workbook()
    parsed = xs.parse_workbook(buf)
    for _, title, _ in xs.ENTITY_SHEETS.values():
        assert title in parsed
        assert parsed[title] == []  # blank template: headers only, no data rows
    for sheet_name in xs.MONTHLY_SHEETS:
        assert sheet_name in parsed


def test_sample_workbook_round_trips_through_validation_with_zero_errors():
    buf = xs.generate_sample_workbook()
    parsed = xs.parse_workbook(buf)
    issues = xs.validate_import(parsed)
    errors = [i for i in issues if i["level"] == "error"]
    assert errors == [], f"Sample dataset must always pass validation, got: {errors}"
    # the sample is built to exercise every sheet, not just master data
    assert len(parsed["Sources"]) >= 2
    assert len(parsed["Monthly_Supply"]) > 0
    assert len(parsed["Monthly_Demand"]) > 0


def test_sample_workbook_imports_successfully_via_client(client):
    r = client.post("/api/projects", json={"name": "Excel Import Target"})
    pid = r.get_json()["id"]

    sample_buf = xs.generate_sample_workbook()
    r = client.post(f"/api/data/{pid}/validate", data={"file": (io.BytesIO(sample_buf.getvalue()), "sample.xlsx")},
                     content_type="multipart/form-data")
    assert r.status_code == 200
    preview = r.get_json()
    assert preview["status"] == "PASSED"
    assert preview["error_count"] == 0

    r = client.post(f"/api/data/{pid}/import", data={"file": (io.BytesIO(sample_buf.getvalue()), "sample.xlsx")},
                     content_type="multipart/form-data")
    assert r.status_code == 200
    result = r.get_json()
    assert result["status"] == "READY FOR OPTIMIZATION"
    assert result["added"]["sources"] == 2
    assert result["monthly_records_imported"] > 0

    # imported data should actually be usable by the (single-period AND
    # multi-period) optimizer - not just sitting inert in the project file
    r = client.post(f"/api/optimize/{pid}/multiperiod", json={})
    res = r.get_json()
    assert res["status"] == "optimal"
    assert len(res["monthly_breakdown"]) == 12


def test_validation_catches_bad_reference_negative_value_and_duplicate_month():
    parsed = xs.parse_workbook(xs.generate_sample_workbook())

    # bad reference: corridor points at a source that doesn't exist
    parsed["Corridors"][0]["origin_code"] = "SRC-DOES-NOT-EXIST"
    # negative capacity on a CGS row
    parsed["CGS"][0]["capacity"] = -100
    # duplicate (source_id, month) row in Monthly_Supply
    dup = dict(parsed["Monthly_Supply"][0])
    parsed["Monthly_Supply"].append(dup)

    issues = xs.validate_import(parsed)
    messages = [i["message"] for i in issues if i["level"] == "error"]
    assert any("not a known Sources.code" in m for m in messages)
    assert any("capacity" in m and ">= 0" in m for m in messages)
    assert any("Duplicate row" in m for m in messages)


def test_missing_required_sheet_reported_as_schema_error():
    parsed = xs.parse_workbook(xs.generate_sample_workbook())
    del parsed["Priority_Tiers"]
    issues = xs.validate_import(parsed)
    assert any(i["level"] == "error" and "Priority_Tiers" in i["message"] for i in issues)


def test_import_with_errors_is_rejected_and_writes_nothing(client):
    r = client.post("/api/projects", json={"name": "Reject Bad Import"})
    pid = r.get_json()["id"]

    from openpyxl import load_workbook
    wb = load_workbook(xs.generate_sample_workbook())
    ws = wb["Corridors"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    origin_col = header.index("origin_code") + 1
    ws.cell(row=2, column=origin_col).value = "NOT-A-REAL-SOURCE"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    r = client.post(f"/api/data/{pid}/import", data={"file": (out, "broken.xlsx")},
                     content_type="multipart/form-data")
    assert r.status_code == 400
    assert r.get_json()["error_count"] > 0

    r = client.get(f"/api/network/{pid}/sources")
    assert r.get_json() == []  # nothing was written
